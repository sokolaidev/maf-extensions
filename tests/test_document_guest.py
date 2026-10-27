"""Real guest conversions and hostile input guards, without a backend or network."""

from __future__ import annotations

import io
import json
import re
import shutil
import struct
import subprocess
import sys
import zipfile
import zlib
from importlib.metadata import version
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from document_fixtures import documents, docx, pdf, pdf_objects
from maf_sandbox_document import _guest
from maf_sandbox_document._pdf_guard import Refused, check_pdf

SOURCE = Path(_guest.__file__).parent
DEFAULTS = dict(
    kind="pdf",
    format="text",
    pages=None,
    max_pages=200,
    max_chars=1_000_000,
    image_dpi=96,
    max_image_pixels=16_000_000,
    max_image_pages=20,
)


def test_guest_tests_use_the_image_dependency_versions():
    manifest = SOURCE.parents[3] / "images/document-sandbox/requirements.txt"
    pins = re.findall(r"^([a-zA-Z0-9_-]+)==([0-9.]+)", manifest.read_text(), re.MULTILINE)
    assert len(pins) == 9
    assert all(version(name) == expected for name, expected in pins)


def run_guest(directory: Path, data: bytes, **options):
    directory.mkdir(parents=True, exist_ok=True)
    for name in ("_guest.py", "_pdf_guard.py"):
        shutil.copyfile(SOURCE / name, directory / name)
    (directory / "input.bin").write_bytes(data)
    (directory / "request.json").write_text(json.dumps(DEFAULTS | options))
    result = subprocess.run(
        [sys.executable, "-I", "_guest.py"],
        cwd=directory,
        capture_output=True,
        text=True,
        timeout=20,
    )
    metadata = directory / "document.json"
    assert metadata.exists(), (result.returncode, result.stdout, result.stderr)
    return result, json.loads(metadata.read_text())


@pytest.mark.parametrize("kind", ["pdf", "docx", "pptx", "rtf"])
@pytest.mark.parametrize("format", ["text", "markdown"])
def test_real_formats(tmp_path, kind, format):
    result, metadata = run_guest(tmp_path, documents()[f"fixture.{kind}"], kind=kind, format=format)
    assert result.returncode == 0, result.stderr
    name = "document.md" if format == "markdown" else "document.txt"
    text = (tmp_path / name).read_text("utf-8")
    assert "Document fixture" in text
    if kind == "pptx":
        assert "table cell" in text and "Notes 1" in text and "Notes 2" in text
    assert metadata == {
        "schema_version": 1,
        "media_type": _guest.MEDIA_TYPES[kind],
        "format": format,
        "supported_formats": ["text", "markdown"] + (["images"] if kind == "pdf" else []),
        "pages": 2 if kind in ("pdf", "pptx") else None,
        "selected_pages": [1, 2] if kind in ("pdf", "pptx") else [],
        "characters": len(text),
        "truncated": False,
        "range_clamped": False,
        "verdict": "extracted",
        "reason": None,
        "files": [name],
    }


def test_pdf_images_are_bounded_pngs_with_original_page_numbers(tmp_path):
    result, metadata = run_guest(tmp_path, pdf(), format="images", pages=[2, 100])
    assert result.returncode == 0, result.stderr
    assert metadata["files"] == ["page-2.png"] and metadata["range_clamped"]
    raw = (tmp_path / "page-2.png").read_bytes()
    assert raw[:8] == b"\x89PNG\r\n\x1a\n"
    assert struct.unpack(">II", raw[16:24]) == (192, 96)


@pytest.mark.parametrize(
    "options,cap",
    [
        ({"max_pages": 1}, "max_pages"),
        ({"max_chars": 4}, "max_chars"),
        ({"format": "images", "max_image_pixels": 100}, "max_image_pixels"),
        ({"format": "images", "max_image_pages": 1}, "max_image_pages"),
    ],
)
def test_pdf_caps(tmp_path, options, cap):
    result, metadata = run_guest(tmp_path, pdf(), **options)
    assert result.returncode == 2
    assert metadata["verdict"] == "refused" and metadata["files"] == []
    assert cap in metadata["reason"]
    assert not metadata["truncated"]


def test_rtf_characters_and_ignored_page_range(tmp_path):
    result, metadata = run_guest(
        tmp_path, rb"{\rtf1 hello}", kind="rtf", pages=[5, 10], max_chars=4
    )
    assert result.returncode == 2 and "max_chars" in metadata["reason"]
    result, metadata = run_guest(tmp_path, rb"{\rtf1 hello}", kind="rtf", pages=[5, 10])
    assert result.returncode == 0 and metadata["pages"] is None


@pytest.mark.parametrize("kind", ["docx", "pptx", "rtf"])
def test_non_pdf_images_are_explicitly_refused(tmp_path, kind):
    result, metadata = run_guest(
        tmp_path, documents()[f"fixture.{kind}"], kind=kind, format="images"
    )
    assert result.returncode == 2 and "not supported" in metadata["reason"]
    assert metadata["supported_formats"] == ["text", "markdown"]


def archive_with(name, data):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(name, data)
    return output.getvalue().replace(name.replace(chr(92), "/").encode(), name.encode())


@pytest.mark.parametrize(
    "name,data,reason",
    [
        ("../word/document.xml", b"ok", "unsafe ZIP"),
        ("/word/document.xml", b"ok", "unsafe ZIP"),
        ("word\\document.xml", b"ok", "unsafe ZIP"),
        ("embedded.zip", b"abc", "nested archive"),
        ("embedded.bin", b"PK\x03\x04junk", "nested archive"),
        ("word/document.xml", b"a" * 100_000, "max_zip_ratio"),
        ("word/document.xml", b'<!DOCTYPE a [<!ENTITY b "boom">]><a>&b;</a>', "DTDs"),
        (
            "word/document.xml",
            '<?xml version="1.0" encoding="utf-16"?><!DOCTYPE a [<!ENTITY b "boom">]><a>&b;</a>'.encode(
                "utf-16"
            ),
            "DTDs",
        ),
    ],
    ids=[
        "traversal",
        "absolute",
        "backslash",
        "nested-name",
        "nested-bytes",
        "ratio",
        "entity",
        "utf16-entity",
    ],
)
def test_zip_and_xml_guards(tmp_path, name, data, reason):
    result, metadata = run_guest(tmp_path, archive_with(name, data), kind="docx")
    assert result.returncode == 2 and reason in metadata["reason"]


def test_declared_two_gib_zip_member_is_refused_before_read(tmp_path):
    raw = bytearray(archive_with("word/document.xml", b"<document/>"))
    central = raw.index(b"PK\x01\x02")
    struct.pack_into("<I", raw, central + 24, 2 * 1024**3)
    result, metadata = run_guest(tmp_path, raw, kind="docx")
    assert result.returncode == 2 and "max_zip_member_bytes" in metadata["reason"]


def test_truncated_ooxml(tmp_path):
    result, metadata = run_guest(tmp_path, docx()[:100], kind="docx")
    assert result.returncode == 2 and "invalid document" in metadata["reason"]


def test_pdf_stream_expansion_refused_before_pdfium(tmp_path):
    # A low cap keeps the fixture small while exercising the actual bounded decompressor.
    from maf_sandbox_document import _pdf_guard

    original = _pdf_guard.MAX_STREAM_BYTES
    _pdf_guard.MAX_STREAM_BYTES = 1024
    try:
        with pytest.raises(Refused, match="max_pdf_stream_bytes"):
            check_pdf(pdf(text="x" * 2000))
    finally:
        _pdf_guard.MAX_STREAM_BYTES = original


@pytest.mark.parametrize("pages,chosen", [(None, [1, 2]), ([2, 99], [2]), ([9, 99], [2])])
def test_range_clamping(tmp_path, pages, chosen):
    result, metadata = run_guest(tmp_path, pdf(), pages=pages)
    assert result.returncode == 0 and metadata["selected_pages"] == chosen


def test_docx_range_refused(tmp_path):
    result, metadata = run_guest(tmp_path, docx(), kind="docx", pages=[1, 2])
    assert result.returncode == 2 and "pagination engine" in metadata["reason"]


@pytest.mark.parametrize(
    "guard,limit,reason",
    [
        ("MAX_MEMBERS", 1, "max_zip_members"),
        ("MAX_TOTAL_BYTES", 1, "max_zip_total_bytes"),
    ],
)
def test_zip_count_and_total_guards(monkeypatch, guard, limit, reason):
    monkeypatch.setattr(_guest, guard, limit)
    with pytest.raises(_guest.Refused, match=reason):
        _guest.guarded_zip(docx(), "docx")


def test_total_pdf_stream_limit(monkeypatch):
    from maf_sandbox_document import _pdf_guard

    monkeypatch.setattr(_pdf_guard, "MAX_STREAM_TOTAL", 1)
    with pytest.raises(Refused, match="max_pdf_stream_total"):
        check_pdf(pdf())


def test_unknown_pdf_filter_is_not_passed_to_pdfium():
    with pytest.raises(Refused, match="unsupported PDF stream filter"):
        check_pdf(pdf().replace(b"/FlateDecode", b"/BogusDecode"))


def test_unresolved_pdf_reference_is_not_repaired():
    with pytest.raises(Refused, match="unresolved PDF reference"):
        check_pdf(pdf().replace(b"/Root 1 0 R", b"/Root 9 0 R"))


def test_incremental_pdf_is_refused():
    with pytest.raises(Refused, match="unsupported PDF trailer"):
        check_pdf(pdf().replace(b"/Size ", b"/Prev 0 /Size "))


def test_inline_image_detection_distinguishes_text_and_comments():
    from maf_sandbox_document._pdf_guard import _has_inline_image

    assert _has_inline_image(b"q BI/W 10/H 10 ID ... EI Q")
    assert _has_inline_image(b"q 1BI /W 10 ID ... EI")
    assert not _has_inline_image(b"BT (BI report) Tj ET")
    assert not _has_inline_image(b"% BI\nBT <4249> Tj ET")
    assert not _has_inline_image(b"/BI /Tag BDC EMC")
    check_pdf(pdf(text="BI report"))


def test_xml_in_nonstandard_member_name_is_guarded():
    raw = archive_with("custom.part", b'<!DOCTYPE a [<!ENTITY b "boom">]><a>&b;</a>')
    with pytest.raises(_guest.Refused, match="DTDs"):
        _guest.guarded_zip(raw, "docx")


def test_duplicate_zip_names_are_refused():
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("a.xml", b"<a/>")
        with pytest.warns(UserWarning, match="Duplicate"):
            archive.writestr("a.xml", b"<b/>")
    with pytest.raises(_guest.Refused, match="unsafe ZIP"):
        _guest.guarded_zip(stream.getvalue(), "docx")


def test_explicit_billion_laughs_is_refused(tmp_path):
    entities = '<!ENTITY a "ha">' + "".join(
        f'<!ENTITY a{i} "' + (("&a;" if i == 1 else f"&a{i - 1};") * 10) + '">'
        for i in range(1, 10)
    )
    raw = archive_with("word/document.xml", f"<!DOCTYPE a [{entities}]><a>&a9;</a>".encode())
    result, metadata = run_guest(tmp_path, raw, kind="docx")
    assert result.returncode == 2 and "DTDs" in metadata["reason"]


def presentation_with_groups(*, missing_notes=False):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, Inches(1), Inches(1)).text = "first grouped text"
    nested = group.shapes.add_group_shape()
    nested.shapes.add_textbox(0, 0, Inches(1), Inches(1)).text = "nested grouped text"
    table = slide.shapes.add_table(1, 1, 0, 0, Inches(1), Inches(1))
    table.table.cell(0, 0).text = "grouped table cell"
    nested.shapes._spTree.insert_element_before(table._element, "p:extLst")
    slide.shapes.add_textbox(0, 0, Inches(1), Inches(1)).text = "last top-level text"
    notes = slide.notes_slide.notes_text_frame
    assert notes is not None
    notes.text = "notes text"
    if missing_notes:
        for shape in list(slide.notes_slide.shapes):
            element = shape._element
            element.getparent().remove(element)
        assert slide.notes_slide.notes_text_frame is None
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


@pytest.mark.parametrize("format", ["text", "markdown"])
@pytest.mark.parametrize("missing_notes", [False, True])
def test_pptx_nested_groups_and_optional_notes(tmp_path, format, missing_notes):
    result, metadata = run_guest(
        tmp_path, presentation_with_groups(missing_notes=missing_notes), kind="pptx", format=format
    )
    assert result.returncode == 0, result.stderr
    text = (tmp_path / metadata["files"][0]).read_text("utf-8")
    ordered = [
        "first grouped text",
        "nested grouped text",
        "grouped table cell",
        "last top-level text",
    ]
    assert all(text.count(value) == 1 for value in ordered)
    assert [text.index(value) for value in ordered] == sorted(
        text.index(value) for value in ordered
    )
    assert ("notes text" in text) is not missing_notes
    assert metadata["characters"] == len(text) and metadata["truncated"] is False


def test_pptx_grouped_text_keeps_character_limit(tmp_path):
    result, metadata = run_guest(tmp_path, presentation_with_groups(), kind="pptx", max_chars=50)
    assert result.returncode == 2 and "max_chars" in metadata["reason"]


def pdf_stream(data, attributes=b""):
    encoded = zlib.compress(data)
    return (
        b"<< /Filter /FlateDecode /Length "
        + str(len(encoded)).encode()
        + b" "
        + attributes
        + b" >>\nstream\n"
        + encoded
        + b"\nendstream"
    )


def image_pdf(pixels, *, content_reference=b"4 0 R"):
    return pdf_objects(
        [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 10 10] /Resources "
            b"<< /XObject << /Im 5 0 R >> >> /Contents " + content_reference + b" >>",
            pdf_stream(b"q 10 0 0 10 0 0 cm /Im Do Q"),
            pdf_stream(
                pixels,
                b"/Type /XObject /Subtype /Image /Width 1 /Height 1 "
                b"/ColorSpace /DeviceRGB /BitsPerComponent 8",
            ),
        ]
    )


@pytest.mark.parametrize("pixels", [b"ABC", b"BI ", b"(\x00\x00"])
def test_pdf_image_bytes_are_not_drawing_syntax(tmp_path, pixels):
    raw = image_pdf(pixels)
    check_pdf(raw)
    result, metadata = run_guest(tmp_path, raw, format="images")
    assert result.returncode == 0, result.stderr
    assert metadata["files"] == ["page-1.png"]
    assert (tmp_path / "page-1.png").read_bytes().startswith(b"\x89PNG\r\n\x1a\n")


def test_pdf_image_reused_as_page_content_still_checks_drawing_syntax():
    with pytest.raises(Refused, match="inline PDF images"):
        check_pdf(image_pdf(b"BI ", content_reference=b"5 0 R"))


@pytest.mark.parametrize(
    "usage", ["form", "pattern", "charproc", "direct-charproc", "split-page", "indirect-page"]
)
def test_pdf_executable_streams_still_refuse_inline_images(usage):
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 10 10] /Contents 4 0 R >>",
        pdf_stream(b"q Q"),
    ]
    if usage == "form":
        objects.append(pdf_stream(b"BI /W 1 ID x EI", b"/Subtype /Form"))
    elif usage == "pattern":
        objects.append(pdf_stream(b"BI /W 1 ID x EI", b"/PatternType 1"))
    elif usage in ("charproc", "direct-charproc"):
        objects.extend(
            [
                b"<< /Type /Font /Subtype /Type3 /CharProcs 6 0 R >>",
                b"<< /A 7 0 R >>",
                pdf_stream(b"BI /W 1 ID x EI"),
            ]
        )
        if usage == "direct-charproc":
            objects[4] = b"<< /Font << /F1 " + objects[4] + b" >> >>"
    else:
        objects[3] = b"[5 0 R 6 0 R]" if usage == "indirect-page" else pdf_stream(b"B")
        objects.extend([pdf_stream(b"B"), pdf_stream(b"I /W 1 ID x EI")])
        if usage == "split-page":
            objects[2] = objects[2].replace(b"4 0 R", b"[5 0 R 6 0 R]")
    with pytest.raises(Refused, match="inline PDF images"):
        check_pdf(pdf_objects(objects))


def test_pdf_binary_streams_keep_decompression_limit(monkeypatch):
    from maf_sandbox_document import _pdf_guard

    monkeypatch.setattr(_pdf_guard, "MAX_STREAM_BYTES", 100)
    with pytest.raises(Refused, match="max_pdf_stream_bytes"):
        check_pdf(image_pdf(b"X" * 200))


def test_pdf_font_file_bytes_are_not_drawing_syntax():
    check_pdf(
        pdf_objects(
            [
                b"<< /Type /Catalog /Pages 2 0 R >>",
                b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
                b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 10 10] /Contents 4 0 R >>",
                pdf_stream(b"q Q"),
                b"<< /Type /FontDescriptor /FontFile2 6 0 R >>",
                pdf_stream(b"BI (\x00\x00"),
            ]
        )
    )


def test_pdf_repeated_content_keeps_total_scan_bound(monkeypatch):
    from maf_sandbox_document import _pdf_guard

    monkeypatch.setattr(_pdf_guard, "MAX_STREAM_TOTAL", 100)
    references = b"[" + b"4 0 R " * 10 + b"]"
    with pytest.raises(Refused, match="max_pdf_content_total"):
        check_pdf(image_pdf(b"ABC", content_reference=references))


def test_pdf_content_reference_cycles_are_refused():
    with pytest.raises(Refused, match="max_pdf_structure"):
        check_pdf(
            pdf_objects(
                [
                    b"<< /Type /Catalog /Pages 2 0 R >>",
                    b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
                    b"<< /Type /Page /Parent 2 0 R /Contents 4 0 R >>",
                    b"[4 0 R]",
                ]
            )
        )
