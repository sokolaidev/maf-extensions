"""Small generated documents for offline and backend qualification."""

from __future__ import annotations

import io
import zipfile
import zlib


def pdf(pages: int = 2, text: str = "Document fixture", *, compressed: bool = True) -> bytes:
    """Build a classic PDF with a selectable number of real text pages."""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids ["
        + b" ".join(f"{4 + 2 * i} 0 R".encode() for i in range(pages))
        + f"] /Count {pages} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for index in range(pages):
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 144 72] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * index} 0 R >>".encode()
        )
        content = f"BT /F1 12 Tf 10 40 Td ({text} {index + 1}) Tj ET".encode()
        if compressed:
            content = zlib.compress(content)
        objects.append(
            f"<< /Length {len(content)}".encode()
            + (b" /Filter /FlateDecode" if compressed else b"")
            + b" >>\nstream\n"
            + content
            + b"\nendstream"
        )
    data = bytearray(b"%PDF-1.7\n")
    offsets = []
    for number, value in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{number} 0 obj\n".encode() + value + b"\nendobj\n")
    crossref = len(data)
    data.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(
        f"trailer\n<< /Root 1 0 R /Size {len(objects) + 1} >>\n"
        f"startxref\n{crossref}\n%%EOF\n".encode()
    )
    return bytes(data)


def docx(text: str = "Document fixture") -> bytes:
    """Build a minimal Word document with a heading and body paragraph."""
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            "</Types>",
        )
        archive.writestr(
            "_rels/.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            "</Relationships>",
        )
        archive.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:body><w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr>'
            f"<w:r><w:t>{text}</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>Second paragraph</w:t></w:r></w:p></w:body></w:document>",
        )
    return output.getvalue()


def pptx() -> bytes:
    """Build two slides, a table and notes using the pinned fixture dependency."""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    for number in (1, 2):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        title = slide.shapes.title
        assert title is not None
        title.text = f"Document fixture {number}"
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(3), Inches(1)).table
        table.cell(0, 0).text = "table cell"
        table.cell(0, 1).text = "value"
        notes = slide.notes_slide.notes_text_frame
        assert notes is not None
        notes.text = f"Notes {number}"
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def documents() -> dict[str, bytes]:
    """One real fixture for each supported document type."""
    return {
        "fixture.pdf": pdf(),
        "fixture.docx": docx(),
        "fixture.pptx": pptx(),
        "fixture.rtf": rb"{\rtf1\ansi Document fixture\par Second paragraph}",
    }
